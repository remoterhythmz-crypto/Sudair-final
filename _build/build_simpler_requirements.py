from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor

R=Path(__file__).resolve().parent
md=R/'Sudair_Client_Data_Requirements_Simple_v2.md'
doc=R/'Sudair_Client_Data_Requirements_Simple_v2.docx'
known=[
 ('Site footprint','72,000 m² site; 12,000 m² built area','Operations page'),
 ('Oncology areas','2,826.9 m² production; 1,848.6 m² storage; 70.3 m² refrigerated storage (2–8°C)','Projects page'),
 ('Oncology line capacity','300M tablets; 150M capsules; 9M sterile liquid units; 6M lyophilised vials, stated for one shift/day','Projects page'),
 ('Projects','Insulin, biologics and respiratory facilities described in the roadmap','Projects page'),
]
asks=[
 ('1. Current site and departments','A dated map showing each operating, commissioning and planned building; current area for production, QC, QA, warehousing, utilities and offices. Reconcile the 72,000 m² website figure with any other figure used in presentations.'),
 ('2. Lines and capacity','A one-page line table: dosage form, line/equipment name, installed and qualified status, rated capacity with time/shift denominator, latest actual output period, utilisation, and whether each number may be public.'),
 ('3. Quality and technology','What QC, QA and validation teams do; approved photos of their actual spaces; the exact automation, containment, isolator, monitoring or digital systems in use, with vendor/name and public-use permission. Separate operating technology from planned technology.'),
 ('4. People and partners','Names and roles of department heads, years of relevant experience, prior affiliations or qualifications that may be published, approved short quotes, portrait consent and partner-name/logo permission.'),
 ('5. Image identification and rights','For each selected photo or footage file: facility, department, date, people depicted, photographer/owner, written release, and any restricted equipment, label, screen or third-party mark. Identify which first-wave images are illustrative and must not be presented as documentary Sudair photographs.'),
 ('6. Sign-off and brand','Current master logo and brand guide, the approver for corporate facts and medical wording, and which facility/capacity/leadership facts are cleared for public social use.'),
]
lines=['# Sudair client data request — concise version','',
'**Purpose:** give the three-month corporate campaign specific, current and publishable evidence. Please return a fact sheet and matching photo/footage IDs, with an owner and public-use status for each item.','',
'## Already published on Sudair’s website — confirm currency before reuse','',
'| Topic | Published statement | Source |','|---|---|---|']
for a,b,c in known:lines.append(f'| {a} | {b} | {c} |')
lines+=['',
'The site-area figures and oncology-area figures describe different scopes and must not be combined. The capacity numbers are **one-shift stated capacity**, not actual output. The website describes several facilities as projects; request the operating status of each before calling it operational.','',
'[Operations page](https://sudairpharma.com/operations/) · [Projects page](https://sudairpharma.com/projects/)','',
'## Please provide','']
for a,b in asks:lines += [f'### {a}','',b,'']
lines += ['## Minimum first response','',
'1. Current facility/department map and operating-status legend.','2. Line-by-line capacity sheet with denominator and latest actual-output period.','3. QC/QA and technology fact sheet with approved images.','4. Leadership list with public bios and consent.','5. Asset-identification and rights register for the images used in the proposed posts.','']
md.write_text('\n'.join(lines),encoding='utf-8')

d=Document();sec=d.sections[0];sec.top_margin=Inches(.42);sec.bottom_margin=Inches(.42)
sec.left_margin=sec.right_margin=Inches(.62)
styles=d.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(8)
styles['Normal'].paragraph_format.space_after=Pt(2)
for n in ['Heading 1','Heading 2']:
    styles[n].font.name='Aptos Display';styles[n].font.color.rgb=RGBColor(0,99,68)
    styles[n].paragraph_format.space_before=Pt(5)
    styles[n].paragraph_format.space_after=Pt(2)
styles['Heading 2'].font.size=Pt(10)
styles['Title'].font.size=Pt(18)
d.add_heading('Sudair: information needed for the 3-month campaign',0)
d.add_paragraph('Please confirm current, publicly usable facts and identify the matching source photos. A short fact sheet is enough; the questions below focus on gaps in the material supplied.')
d.add_heading('What the website already says',1)
t=d.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,h in zip(t.rows[0].cells,['Topic','Published figure / status','Source']):c.text=h
for row in known:
    cells=t.add_row().cells
    for c,v in zip(cells,row):c.text=v
d.add_paragraph('These figures describe different scopes. Capacity is stated for one shift, not actual output; planned projects must not be described as operating without confirmation.')
d.add_paragraph('Sources: sudairpharma.com/operations/ and sudairpharma.com/projects/')
d.add_heading('Please send six short answers',1)
for a,b in asks:
    d.add_heading(a,2);d.add_paragraph(b)
d.add_paragraph('First return: facility map, capacity sheet, QC/QA and technology notes, public leadership bios, and a rights register. Mark each item approved for public use, internal only, or not confirmed.')
d.save(doc)
print(md);print(doc)
