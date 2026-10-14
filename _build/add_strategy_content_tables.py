# Add the strategy + content tables from the "Last Version" 6-month plan (V2-8)
# into the 3-month strategy template (first-wave image version, v2).
# Output: a NEW file in Sudair-final. Nothing else is touched.
import shutil
from pathlib import Path

import lxml.etree as etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PACK = Path(r"D:\Medical communications- 2025\Sudair\Corporate_Sudair_Social\Sudair_3Month_Pack")
BASE = PACK / "Sudair_3Month_Strategy_First_Wave_Template_v2.pptx"
SRC = Path(r"D:\Medical communications- 2025\Sudair\Corporate_Sudair_Social\Last Version\Sudair_Content_Strategy_and_6Month_Plan V2-8.pptx")
OUT = PACK / "Sudair-final" / "Sudair_3Month_Strategy_First_Wave_Tables_v3.pptx"

BG = "F7F9F6"
BAR = "09A45B"
TITLE_C = "0B4F35"
SUB_C = "5B6B63"
HDR_FILL = "0B4F35"
ROW_A, ROW_B = "FFFFFF", "F1F5F1"
TXT = "1E2B24"
BORDER = "D9E2DC"
FONT = "Segoe UI"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- extract ----
def extract_tables():
    src = Presentation(str(SRC))
    out = []
    for slide in src.slides:
        title = ""
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                title = sh.text_frame.text.strip().replace("\n", " / ")
                break
        for sh in slide.shapes:
            if getattr(sh, "has_table", False):
                t = sh.table
                data = [[cell.text for cell in row.cells] for row in t.rows]
                widths = [col.width for col in t.columns]
                out.append({"title": title, "data": data, "widths": widths})
    return out


# ---------------------------------------------------------------- helpers ----
def add_bg(slide, color=BG):
    xml = (
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<p:bgPr><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'
    )
    cSld = slide._element.find(qn("p:cSld"))
    cSld.insert(0, etree.fromstring(xml))


def add_topbar(slide):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.34), Inches(0.1))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(BAR)
    shp.line.fill.background()
    shp.shadow.inherit = False


def add_text(slide, text, left, top, width, height, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def set_cell_borders(cell, color=BORDER, w=6350):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):  # inserted at 0 -> final order L,R,T,B
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(w))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", color)
        dash = etree.SubElement(ln, qn("a:prstDash"))
        dash.set("val", "solid")
        tcPr.remove(ln)
        tcPr.insert(0, ln)


def fill_cell(cell, text, size, bold, color, fill_hex, anchor=MSO_ANCHOR.TOP):
    cell.vertical_anchor = anchor
    cell.margin_left = Inches(0.055)
    cell.margin_right = Inches(0.055)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill_hex)
    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    lines = [ln for ln in (text or "").replace("\v", "\n").split("\n")]
    if not lines:
        lines = [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT if False else PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        p.space_after = Pt(0)
        p.space_before = Pt(0)
    set_cell_borders(cell)


def pick_font(data):
    rows = len(data)
    cols = len(data[0]) if data else 1
    longest = max((len(str(c)) for row in data for c in row), default=0)
    if rows >= 14:
        return 6.5
    if longest > 170 or (rows >= 9 and longest > 90):
        return 7.5
    if longest > 95 or rows >= 9:
        return 8.0
    return 9.0


def build_table_slide(prs, spec, index, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_topbar(slide)
    add_text(slide, spec["title"], 0.5, 0.24, 11.6, 0.6, 22, TITLE_C, bold=True)
    add_text(
        slide,
        "Strategy and content tables · imported from Content Strategy and Six-Month Plan V2-8",
        0.5, 0.80, 11.6, 0.35, 11, SUB_C,
    )
    add_text(slide, f"tables {index}/{total}", 12.0, 7.10, 1.0, 0.3, 9, SUB_C, align=PP_ALIGN.RIGHT)

    data = spec["data"]
    nrows, ncols = len(data), len(data[0])
    font = pick_font(data)

    left, top, width = Inches(0.5), Inches(1.32), Inches(12.33)
    height = Inches(5.7)
    gf = slide.shapes.add_table(nrows, ncols, left, top, width, height)
    tbl = gf.table

    # no default blue style
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for style in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(style)
    sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid

    # column widths scaled from the source design
    src_w = spec["widths"]
    tot = sum(src_w) or 1
    scale = 12.33 / (tot / 914400.0)
    for c, w in enumerate(src_w):
        tbl.columns[c].width = Emu(int(w * scale))

    # row heights
    hdr_h = Inches(0.34)
    body_h = Inches(max(0.26, (5.7 - 0.34) / max(1, nrows - 1)))
    for r in range(nrows):
        tbl.rows[r].height = hdr_h if r == 0 else Emu(int(body_h))

    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            if r == 0:
                fill_cell(cell, str(val), min(font, 8.5), True, "FFFFFF", HDR_FILL, MSO_ANCHOR.MIDDLE)
            else:
                fill_cell(
                    cell, str(val), font, False, TXT,
                    ROW_A if r % 2 == 1 else ROW_B, MSO_ANCHOR.TOP,
                )
    return slide


def build_divider(prs, specs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_topbar(slide)
    add_text(slide, "Strategy and content tables", 0.5, 1.6, 12.3, 1.0, 34, TITLE_C, bold=True)
    add_text(
        slide,
        "Reference tables imported from Sudair_Content_Strategy_and_6Month_Plan V2-8 "
        "(the last provided strategy file). Strategy tables first, content tables after.",
        0.5, 2.5, 12.3, 0.8, 14, SUB_C,
    )
    # two-column list of the tables
    names = [s["title"] for s in specs]
    half = (len(names) + 1) // 2
    for col, chunk in enumerate((names[:half], names[half:])):
        box = slide.shapes.add_textbox(Inches(0.6 + col * 6.3), Inches(3.5), Inches(6.0), Inches(3.2))
        tf = box.text_frame
        tf.word_wrap = True
        for i, name in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = "· " + name
            run.font.name = FONT
            run.font.size = Pt(11)
            run.font.color.rgb = rgb(TXT)
            p.space_after = Pt(6)
    return slide


def main():
    specs = extract_tables()
    print("tables found in source:", len(specs))
    for s in specs:
        print("  -", s["title"], "|", len(s["data"]), "x", len(s["data"][0]))

    tmp = OUT.with_suffix(".tmp.pptx")
    shutil.copy2(BASE, tmp)
    prs = Presentation(str(tmp))
    before = len(prs.slides._sldIdLst)

    # insertion point: before the November month divider (fallback: after slide 12)
    target = 12
    for i, slide in enumerate(prs.slides):
        txt = " ".join(
            sh.text_frame.text for sh in slide.shapes if sh.has_text_frame
        )
        if txt.strip().startswith("November 2026") and "32 pieces" in txt:
            target = i
            break
    print("insert before slide index (0-based):", target)

    divider = build_divider(prs, specs)
    for i, spec in enumerate(specs, 1):
        build_table_slide(prs, spec, i, len(specs))

    # move the appended slides to the insertion point, in order
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    new_ids = ids[before:]
    for k, nid in enumerate(new_ids):
        sldIdLst.remove(nid)
        sldIdLst.insert(target + k, nid)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(str(OUT))
    tmp.unlink(missing_ok=True)

    check = Presentation(str(OUT))
    print("saved:", OUT, "| slides:", len(check.slides._sldIdLst))


if __name__ == "__main__":
    main()
