# -*- coding: utf-8 -*-
"""Build Sudair_Facilities_Data_and_Assets_Request.docx — the missing-information list."""
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"D:\Medical communications- 2025\Sudair\Corporate_Sudair_Social\Sudair_3Month_Pack\Content_Review\Sudair_Facilities_Data_and_Assets_Request.docx"

d = docx.Document()
st = d.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(11)

def h(text, level=1):
    p = d.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x09, 0x3F, 0x30)
    return p

def bullet(text):
    d.add_paragraph(text, style='List Bullet')

title = d.add_heading('Sudair City — Information and Asset Request List', level=0)
for r in title.runs: r.font.color.rgb = RGBColor(0x09, 0x3F, 0x30)

p = d.add_paragraph()
r = p.add_run('Purpose: close the gaps in the social content bank so that future posts can show verified capability — the kind of detail that builds investor and public trust — without holding for facts. '
              'Every item below is either (a) missing data we currently cannot state anywhere, or (b) missing assets we currently have to generate instead of showing the real thing.')
r.font.size = Pt(11)

h('A. Facility footprint and spaces', 1)
bullet('Total site area of Sudair City and built-up area (phased: Phase 1 vs any later phase), in m².')
bullet('Gross floor area per building: production building(s), QC/R&D building, warehouse(s), utilities building, administration.')
bullet('Cleanroom classification and area per zone (ISO 5 / ISO 7 / ISO 8 or GMP Grade A–D), and total cleanroom m².')
bullet('Warehouse types and capacities: ambient, cold room (2–8 °C), freezer (−20 °C), high-potency / contained quarantine areas — each with pallet positions or m².')
bullet('Outdoor utilities footprint: power substation capacity (MVA), HVAC/chiller capacity, WFI / clean steam generation capacity, backup generator capacity, solar plant (if any) MWp.')

h('B. Production lines and capacity', 1)
bullet('Line inventory: dosage form (solid oral, sterile injectable, ointment/cream, liquid, etc.), line name/number, installed year, technology partner, maximum speed (units/hour).')
bullet('Installed capacity vs. actual current output per line (units/year or doses/year), and current utilisation %.')
bullet('High-potency / oncology containment lines: isolator or closed-transfer technology used, occupational exposure band (OEB) level, and what makes them unique in the Kingdom.')
bullet('Insulin line: form (cartridge, pen, vial), technology partner, annual capacity — the molecule that anchors the «made here» story.')
bullet('Serialisation/track-and-trace capability (aggregation, GS1, SFDA LGS integration status).')
bullet('Batch sizes per line (min/max), number of batches per month if disclosable.')

h('C. Quality control laboratory', 1)
bullet('QC lab area, number of analysts, and operating hours (24/7 or day shift).')
bullet('Instrument list at model level: HPLC/UPLC (count and vendor), GC, dissolution, ICP-MS, mass spectrometer, stability chambers (count, capacity, ICH conditions).')
bullet('Stability program scale: number of registered stability protocols running, chambers dedicated to long-term/photostability.')
bullet('Microbiology capability: isolator-based microbiology? media fills count? environmental monitoring programme scope (viable + non-viable, particle counts per zone).')
bullet('Method-development vs. method-verification scope — what QC does in-house vs. sends out.')
bullet('Turnaround time commitment: sample-to-result for release testing (days).')

h('D. Quality assurance systems and certifications', 1)
bullet('Current GMP certificates held: SFDA GMP, EU GMP (which authority), WHO-PQ, any PIC/S reference — with expiry dates.')
bullet('Number of regulatory inspections passed to date and key markets exported to (GCC, MENA, Africa…).')
bullet('QA systems in use: environmental monitoring system, LIMS, MES (manufacturing execution system), paperless batch records? Deviation/CAPA management tool.')
bullet('Any digital innovation in QA: AI-assisted visual inspection, at-line PAT (process analytical technology), continuous process verification pilots.')
bullet('Certifications beyond GMP: ISO 9001, ISO 14001, ISO 45001, ISO 17025 for the lab, energy management (ISO 50001).')

h('E. Innovative and new technology in the facility', 1)
bullet('Any technology first-of-kind in the Kingdom or GCC (e.g. first domestic insulin line, first oncology containment site) — we need the exact claimable wording.')
bullet('Automation level: line automation grade (manual / semi / fully automated), robotic palletising, automated warehouse (AS/RS)?')
bullet('Green building / sustainability features: LEED or Mostadam rating, water recycling %, waste treatment plant capacity, carbon-reduction targets.')
bullet('R&D pipeline: biosimilars in development (molecule + phase), formulation development capability, analytical R&D scope.')
bullet('Technology-transfer alliances (Sanofi, GeroPharm, Dr. Reddy\'s, Accord): what was transferred — equipment, know-how, licence — and what may be shown publicly.')

h('F. Leadership and department heads (for «meet the functions» content)', 1)
bullet('Name, title, department, tenure at Sudair, previous employer/role, and any public affiliation (universities, societies) for: Plant Head, Head of QA, Head of QC, Head of Production, Head of Engineering, Head of Supply Chain, Head of R&D, Head of HR.')
bullet('Which of them may be photographed, quoted or interviewed on camera (consent status per person).')
bullet('Any executive willing to be the recurring LinkedIn voice (newsletter by-line, quarterly video message).')

h('G. Photo and video assets still missing', 1)
bullet('Extracted (unzipped) footage: ~50 GB remains zipped in «Sudair Assets» — extract to allow frame-picking for all 96 planned pieces.')
bullet('Golden-hour / dawn exterior drone pass of the finished facade (current drone files predate landscaping completion).')
bullet('Cleanroom action shots: filling line running, operator through glove ports (sharp, well-lit — the current JAN-F04/S04 frames are blurred).')
bullet('QC lab action shots: analyst at HPLC, sample prep, stability chamber check — the current QC post reuses the lab-team frame.')
bullet('Packaging hall full-team photo (the jaradi group clip is partial).')
bullet('Night-shift set: exterior at night, corridor shift-change, control-room at night.')
bullet('People portraits: 8–10 individual portraits (mixed gender, functions) for the recurring «People of Sudair» card.')
bullet('Event photography: Global Health Exhibition 2025/2026 full set (booth, signing, delegation visits) at original resolution.')

h('H. Data that must stay internal (do not publish without written approval)', 1)
bullet('Any capacity, utilisation, export or inspection numbers → investor-materials channel only, never social.')
bullet('Partner names and molecules → case-by-case written approval from the partner required first.')
bullet('Staff full names → only with signed image consent; otherwise first name + function only.')

p = d.add_paragraph()
r = p.add_run('\nRequested by: Sudair social media working team · For answers, return per section (A–G) in any format; we will restate each answer as approved social copy for review.')
r.font.size = Pt(10); r.font.italic = True

d.save(OUT)
print('saved', OUT)
