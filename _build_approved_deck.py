# -*- coding: utf-8 -*-
"""Build Sudair_3Month_Approved_Images_Deck.pptx — approved stills/stories,
red-marked on slides that need review. Uses Source_Verified images."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image as PILImage

ROOT = r"D:\Medical communications- 2025\Sudair\Corporate_Sudair_Social\Sudair_3Month_Pack"
IMG  = os.path.join(ROOT, "Generated_Stills", "Source_Verified")
OUT  = os.path.join(ROOT, "Content_Review", "Sudair_3Month_Approved_Images_Deck.pptx")

EMERALD = RGBColor(0x09, 0x3F, 0x30)
MINT    = RGBColor(0x8E, 0xF0, 0xB0)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
RED     = RGBColor(0xD9, 0x2B, 0x2B)
AMBER   = RGBColor(0xE8, 0x8A, 0x1A)
GREY    = RGBColor(0x47, 0x5E, 0x54)

# status per image: OK / EDIT / REBUILD (+ optional note)
STATUS = {
 # November
 "NOV-F02": ("OK","Clone out white play-triangle artifact. Consent for staff."),
 "NOV-F04": ("EDIT","Remove stray tanwin under «نقف» → plain «نقف معكم»."),
 "NOV-F06": ("REBUILD","Headline garbled; Ramadan lantern wrong for Nov. Rebuild: real night facade frame + «شكرًا لفريق الليل», no kashida."),
 "NOV-F08": ("REBUILD","3D cartoon off-brand. Rebuild 4-up photographic grid from ALABBASY.8839/8842."),
 "NOV-F09": ("OK","Consent check only."),
 "NOV-F12": ("OK","Composite correct."),
 "NOV-S01": ("OK",""),
 "NOV-S02": ("OK","Add live question sticker at publish."),
 "NOV-S03": ("OK",""),
 "NOV-S04": ("OK",""),
 "NOV-S05": ("EDIT","Cartoon avatars → 3 real cropped portraits from shift clip."),
 "NOV-S06": ("OK",""),
 "NOV-S07": ("OK","Add live sticker."),
 "NOV-S08": ("OK",""),
 "NOV-S09": ("EDIT","Cartoon avatars → real corridor/door frame behind poll."),
 "NOV-S10": ("OK",""),
 "NOV-S11": ("OK","Repeats NOV-F02 frame; acceptable."),
 "NOV-S12": ("OK",""),
 # December
 "DEC-F01": ("OK","Best still of the set."),
 "DEC-F03": ("EDIT","Re-typeset «لفريقنا» (extra hamza rendered); frame repeats NOV-F09."),
 "DEC-F06": ("OK","Consent/usage rights for ~20 faces required."),
 "DEC-F07": ("EDIT","Construction aerial weak for capability msg. Swap hero to finished facade DJI_20260121103516_0070."),
 "DEC-F10": ("REBUILD","Generic AI crystal ball. Replace with real dawn drone frame; keep headline + chip."),
 "DEC-F12": ("OK",""),
 "DEC-S01": ("OK",""),
 "DEC-S02": ("OK",""),
 "DEC-S03": ("OK",""),
 "DEC-S04": ("OK",""),
 "DEC-S05": ("OK","Add live sticker."),
 "DEC-S06": ("OK",""),
 "DEC-S07": ("OK",""),
 "DEC-S08": ("OK",""),
 "DEC-S09": ("OK",""),
 "DEC-S10": ("EDIT","Caption collides with sticker zone — move up 120 px or drop."),
 "DEC-S11": ("OK",""),
 "DEC-S12": ("OK",""),
 # January
 "JAN-F01": ("OK",""),
 "JAN-F03": ("OK",""),
 "JAN-F04": ("REBUILD","Kashida destroyed «الإفراج». Rebuild no-kashida. ⚠ CLAIM HOLD."),
 "JAN-F07": ("EDIT","Blurred frame — re-pull sharp from ALABBASY.8892; fix «الألة»→«الآلة». ⚠ CLAIM HOLD."),
 "JAN-F10": ("REBUILD","Multi-word kashida destroyed headline. Rebuild plain. ⚠ CLAIM HOLD."),
 "JAN-F11": ("EDIT","Photo is lab-team frame mislabelled as quality team — swap frame or retitle."),
 "JAN-S01": ("OK",""),
 "JAN-S02": ("EDIT","Cartoon characters → real line frame jaradi20260121_8608."),
 "JAN-S03": ("OK","Add live sticker."),
 "JAN-S04": ("EDIT","Washed-out hero — re-pull sharper frame ALABBASY.8871. ⚠ CLAIM HOLD."),
 "JAN-S05": ("EDIT","Cartoon poll → type card (T8). ⚠ CLAIM HOLD."),
 "JAN-S06": ("OK",""),
 "JAN-S07": ("OK",""),
 "JAN-S08": ("EDIT","Process statement — keep held until client verifies. No design change. ⚠ CLAIM HOLD."),
 "JAN-S09": ("OK","Add live sticker."),
 "JAN-S10": ("EDIT","Confirm «حصر» → «حصاد»."),
 "JAN-S11": ("EDIT","Third use of same frame — swap to jaradi20260121_8608."),
 "JAN-S12": ("OK",""),
}

TITLES = {
 "NOV-F02": "Meet the team behind the lab","NOV-F04": "World Diabetes Day solidarity","NOV-F06": "Thank you to the night shift",
 "NOV-F08": "Skill in every station","NOV-F09": "People of Sudair: a portrait","NOV-F12": "The month in one image",
 "NOV-S01": "Shift change","NOV-S02": "Question box: life at Sudair","NOV-S03": "Place card: a corner of Sudair City",
 "NOV-S04": "Countdown to 14 Nov","NOV-S05": "Who is on shift now","NOV-S06": "A day in the life",
 "NOV-S07": "Question box: inside the plant","NOV-S08": "Close-up: tools and benches","NOV-S09": "Poll: which part next",
 "NOV-S10": "Countdown: new year at the plant","NOV-S11": "Meet the lab team","NOV-S12": "Place card: facade at golden hour",
 "DEC-F01": "Built here: the site from above","DEC-F03": "Thank you to the teams","DEC-F06": "Thank you to every team",
 "DEC-F07": "Local manufacturing, in one picture","DEC-F10": "After every night, a brighter morning","DEC-F12": "A quiet thank you",
 "DEC-S01": "Good morning from Sudair City","DEC-S02": "Behind the doors","DEC-S03": "Place card: facade among the palms",
 "DEC-S04": "Our year in photos","DEC-S05": "Question box: what next?","DEC-S06": "Place card: from the sky",
 "DEC-S07": "Poll: where is this?","DEC-S08": "A bench, up close","DEC-S09": "Place card: the packaging hall",
 "DEC-S10": "Thank you","DEC-S11": "The night shift","DEC-S12": "See you in January",
 "JAN-F01": "Same site, new year","JAN-F03": "New year, one site","JAN-F04": "Step 6: release",
 "JAN-F07": "Automated, not absent","JAN-F10": "Quality control","JAN-F11": "Meet the quality team",
 "JAN-S01": "Happy new year","JAN-S02": "Week one, same rhythm","JAN-S03": "Question box",
 "JAN-S04": "Step of the day","JAN-S05": "Poll: which step surprised you?","JAN-S06": "Place card of the week",
 "JAN-S07": "Machines and people","JAN-S08": "A validation run, up close","JAN-S09": "Question box",
 "JAN-S10": "Week in review","JAN-S11": "Meet the team","JAN-S12": "Place card of the week",
}

ORDER = ([f"NOV-F{n:02d}" for n in [2,4,6,8,9,12]] + [f"NOV-S{n:02d}" for n in range(1,13)]
       + [f"DEC-F{n:02d}" for n in [1,3,6,7,10,12]] + [f"DEC-S{n:02d}" for n in range(1,13)]
       + [f"JAN-F{n:02d}" for n in [1,3,4,7,10,11]] + [f"JAN-S{n:02d}" for n in range(1,13)])

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color

def tb(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    r.font.name = "Arial"
    return box

# ---- Cover slide
s = prs.slides.add_slide(BLANK); bg(s, EMERALD)
tb(s, Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.4), "Sudair Pharma — 3-Month Approved Images", 40, WHITE, True, PP_ALIGN.CENTER)
tb(s, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.8), "Nov 2026 – Jan 2027 · source-verified stills & stories", 22, MINT, False, PP_ALIGN.CENTER)
tb(s, Inches(0.8), Inches(4.6), Inches(11.7), Inches(1.4),
   "Legend:  ✅ green = approved · 🟠 amber = needs the listed edit · 🔴 red = rebuild before publishing\n"
   "All images internal creative concepts — not cleared publication masters.\n"
   "Consent, Arabic proofread and (January) client regulatory sign-off still required.", 16, WHITE, False, PP_ALIGN.CENTER)

def add_image_slide(pid):
    path = os.path.join(IMG, pid + ".png")
    s = prs.slides.add_slide(BLANK); bg(s, RGBColor(0xF2, 0xF5, 0xF3))
    status, note = STATUS.get(pid, ("OK", ""))
    # left: image
    with PILImage.open(path) as im:
        iw, ih = im.size
    box_h = Inches(6.6); box_w = Inches(4.6)
    ratio = iw / ih
    if ratio > box_w / box_h:
        w = box_w; h = int(box_w / ratio)
    else:
        h = box_h; w = int(box_h * ratio)
    left = Inches(0.5); top = int((SH - h) / 2)
    s.shapes.add_picture(path, left, top, w, h)
    # red/amber frame for flagged
    if status in ("REBUILD", "EDIT"):
        from pptx.enum.shapes import MSO_SHAPE
        frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left - Emu(38100), top - Emu(38100), w + Emu(76200), h + Emu(76200))
        frame.fill.background()
        frame.line.color.rgb = RED if status == "REBUILD" else AMBER
        frame.line.width = Pt(4)
        frame.shadow.inherit = False
    # right: info panel
    lx = Inches(5.5)
    col = {"OK": EMERALD, "EDIT": AMBER, "REBUILD": RED}[status]
    lbl = {"OK": "✅ APPROVED", "EDIT": "🟠 NEEDS EDIT", "REBUILD": "🔴 REBUILD"}[status]
    tb(s, lx, Inches(0.6), Inches(7.2), Inches(0.9), pid, 30, EMERALD, True)
    tb(s, lx, Inches(1.35), Inches(7.2), Inches(0.5), TITLES.get(pid, ""), 18, GREY)
    tb(s, lx, Inches(2.0), Inches(7.2), Inches(0.6), lbl, 22, col, True)
    is_story = "-S" in pid
    tb(s, lx, Inches(2.7), Inches(7.2), Inches(0.5),
       "Channel: Instagram + Facebook  ·  " + ("Story 1080×1920" if is_story else "Feed still 1080×1350"), 14, GREY)
    if note:
        tb(s, lx, Inches(3.3), Inches(7.2), Inches(2.6), "Review note: " + note, 15, RGBColor(0x33,0x33,0x33))
    tb(s, lx, Inches(6.6), Inches(7.2), Inches(0.5),
       "Source set: Generated_Stills/Source_Verified/" + pid + ".png", 11, GREY)

for pid in ORDER:
    if os.path.exists(os.path.join(IMG, pid + ".png")):
        add_image_slide(pid)
    else:
        print("MISSING", pid)

prs.save(OUT)
print("saved", OUT, len(prs.slides.__iter__.__self__._sldIdLst), "slides")
